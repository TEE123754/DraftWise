import io

import openpyxl
import pytest
from docx import Document

from app.ai.grounding import ground_extraction
from app.parsers.registry import parse_document
from app.services.extraction import extract_labelled, identify_role
from app.services.verification import verify

SI = b"""SHIPPING INSTRUCTION
========================================

Shipper: APRIL FAR EAST (M) SDN BHD
  TOWER 2, AVENUE 5, LEVEL 6; KUALA LUMPUR, MALAYSIA
Consignee (Non-Negotiable): EAST BRIGHT FZ-LLC
  RAKEZ AMENITY CENTER; RAK, UAE
Notify: EAST BRIGHT FZ-LLC
Port of Loading (POL): NANTONG, CHINA (CNNTG)
POD: KARACHI, PAKISTAN (PKKHI)
Total Containers: 6 x 40'HC
Gross Wt (kgs): 131,058 KG
"""

BL = b"""BILL OF LADING (DRAFT)
========================================

SHIPPER: APRIL FAR EAST (M) SDN BHD
  TOWER 2, AVENUE 5, LEVEL 6; KUALA LUMPUR, MALAYSIA
To the Order of: UAB NOVAKOPA
  RAKEZ AMENITY CENTER; RAK, UAE
Notify Party: UAB NOVAKOPA
Port of Loading (POL): NANTONG, CHINA (CNNTG)
POD: KARACHI, PAKISTAN (PKKHI)
Container Count: 6 x 40'HC
Gross Weight (KG): 131,058 KG
"""


def extract(content, name):
    document = parse_document(content, name)
    extraction = extract_labelled(document)
    return extraction, ground_extraction(extraction, document)


def test_consignee_and_container_label_variants_are_read():
    si, _ = extract(SI, "si.txt")
    bl, _ = extract(BL, "bl.txt")
    assert si.fields["consignee"].raw_value == "EAST BRIGHT FZ-LLC"
    assert si.fields["container_count"].raw_value == "6 x 40'HC"
    assert bl.fields["consignee"].raw_value == "UAB NOVAKOPA"


def test_label_variants_yield_a_real_mismatch_not_a_vague_review():
    si, si_quality = extract(SI, "si.txt")
    bl, bl_quality = extract(BL, "bl.txt")
    report = verify(si, bl, evidence_quality={**si_quality, **bl_quality})
    assert report.status == "MISMATCH"
    assert {row.field.value for row in report.comparisons if row.decision == "mismatch"} == {
        "consignee",
        "notify_party",
    }


def test_bilingual_word_table_labels_pair_with_the_cell_to_their_right():
    document = Document()
    document.add_paragraph("BILL OF LADING (DRAFT)")
    table = document.add_table(rows=3, cols=2)
    rows = [
        # Chinese written as escapes: 收货人 = consignee, 毛重 = gross weight, 箱数 = container count.
        ("Consignee (\u6536\u8d27\u4eba)", "AL GURG STATIONERY LLC"),
        ("Gross Weight\u6bdb\u91cd(KGS) (\u6bdb\u91cd KGS)", "243,588"),
        ("Total Containers (\u7bb1\u6570)", "12 x 20'FCL"),
    ]
    for row, (label, value) in zip(table.rows, rows, strict=True):
        row.cells[0].text, row.cells[1].text = label, value
    buffer = io.BytesIO()
    document.save(buffer)
    extraction = extract_labelled(parse_document(buffer.getvalue(), "draft.docx"))
    assert extraction.document_type == "BL"
    assert extraction.fields["consignee"].raw_value == "AL GURG STATIONERY LLC"
    assert extraction.fields["gross_weight_kg"].raw_value == "243,588"
    assert extraction.fields["container_count"].raw_value == "12 x 20'FCL"


def test_unfilled_form_fields_are_missing_not_mismatched():
    si_text = SI.replace(b"NANTONG, CHINA (CNNTG)", b"____MT").replace(
        b"KARACHI, PAKISTAN (PKKHI)", b"TBA"
    )
    si, si_quality = extract(si_text, "si.txt")
    assert si.fields["port_of_loading"].state == "missing"
    assert si.fields["port_of_discharge"].state == "missing"
    bl, bl_quality = extract(BL, "bl.txt")
    report = verify(si, bl, evidence_quality={**si_quality, **bl_quality})
    ports = {row.field.value: row.decision for row in report.comparisons if "port" in row.field.value}
    assert ports == {"port_of_loading": "missing", "port_of_discharge": "missing"}


def test_translation_is_stripped_from_labels_but_never_from_values():
    text = "SHIPPING INSTRUCTION\nConsignee (收货人): 上海 ACME CO"
    document = parse_document(text.encode(), "si.txt")
    assert extract_labelled(document).fields["consignee"].raw_value == "上海 ACME CO"


def test_spreadsheet_si_is_recognised_from_its_header_and_sheet_name():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S.I."
    sheet["A1"] = "ACME TRADING PTE LTD"
    sheet["A3"] = "BL INSTRUCTION"
    sheet["A4"], sheet["B4"] = "SHIPPER", "ACME TRADING PTE LTD"
    buffer = io.BytesIO()
    workbook.save(buffer)
    document = parse_document(buffer.getvalue(), "si.xlsx")
    assert identify_role(document) == "SI"
    assert extract_labelled(document).fields["shipper"].raw_value == "ACME TRADING PTE LTD"


def test_ai_output_quoting_a_placeholder_becomes_a_missing_field():
    from app.ai.extraction_contract import Candidate, ExtractionCandidate, expand_candidate
    from app.domain.models import FIELDS

    document = parse_document(SI.replace(b"KARACHI, PAKISTAN (PKKHI)", b"TBA"), "si.txt")
    blocks = {b.text.split(":")[0].strip(): i for i, b in enumerate(document.blocks) if ":" in b.text}
    fields = {name: Candidate() for name in FIELDS}
    fields["port_of_discharge"] = Candidate(value="TBA", blocks=(blocks["POD"],))
    fields["port_of_loading"] = Candidate(
        value="NANTONG, CHINA (CNNTG)", blocks=(blocks["Port of Loading (POL)"],)
    )
    extraction = expand_candidate(ExtractionCandidate(document_type="SI", fields=fields), document)
    assert extraction.fields["port_of_discharge"].state == "missing"
    assert extraction.fields["port_of_loading"].state == "present"


@pytest.mark.parametrize(
    "header",
    ["Booking Instruction", "SI - Shipper's Letter of Instruction", "SI - Shipper\u2019s Letter of Instruction"],
)
def test_instruction_headers_are_recognised_as_si(header):
    document = parse_document(f"{header}\nShipper: ACME TRADING".encode(), "si.txt")
    assert identify_role(document) == "SI"


@pytest.mark.parametrize("header", ["B/L DRAFT FOR APPROVAL", "DRAFT B/L", "BILL OF LADING - DRAFT"])
def test_draft_bl_headers_are_recognised_as_bl(header):
    document = parse_document(f"{header}\nShipper: ACME TRADING".encode(), "bl.txt")
    assert identify_role(document) == "BL"


def test_common_label_synonyms_are_read():
    text = (
        "SHIPPING INSTRUCTIONS\n"
        "Consignor: ACME EXPORTS LTD\nReceiver: BETA IMPORTS LLC\nAlso Notify: GAMMA AGENTS\n"
        "Loading Port: ROTTERDAM\nDestination Port: SANTOS\nQty of Containers: 4 x 40'HC\n"
        "Gross Mass (KG): 21,500 KG\n"
    )
    fields = extract_labelled(parse_document(text.encode(), "si.txt")).fields
    assert {name: fields[name].raw_value for name in fields} == {
        "shipper": "ACME EXPORTS LTD",
        "consignee": "BETA IMPORTS LLC",
        "notify_party": "GAMMA AGENTS",
        "port_of_loading": "ROTTERDAM",
        "port_of_discharge": "SANTOS",
        "container_count": "4 x 40'HC",
        "gross_weight_kg": "21,500 KG",
    }

