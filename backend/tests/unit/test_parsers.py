from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook

from app.ai.grounding import ground_extraction, strict_json
from app.domain.errors import DomainError
from app.domain.models import FieldName
from app.parsers.registry import parse_document
from app.services.extraction import extract_labelled


def test_labels_preserve_evidence_and_reject_package_count():
    parsed = parse_document(
        b"SHIPPING INSTRUCTION\nShipper: Example Export Ltd\nNo. of Containers or Packages: 250 packages\nGross Weight: 22 MT",
        "si",
    )
    result = extract_labelled(parsed)
    assert result.document_type == "SI"
    assert result.fields[FieldName.CONTAINERS].state == "missing"
    assert result.fields[FieldName.WEIGHT].raw_value == "22 MT"
    ground_extraction(result, parsed)


def test_repeated_conflicting_total_is_ambiguous():
    result = extract_labelled(
        parse_document(b"BILL OF LADING\nGross Weight: 22 MT\nGross Weight: 23 MT", "bl")
    )
    assert result.fields[FieldName.WEIGHT].state == "ambiguous"


def test_docx_body_order():
    document = Document()
    document.add_paragraph("SHIPPING INSTRUCTION")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Shipper: Example Export Ltd"
    document.add_paragraph("Gross Weight: 22 MT")
    stream = BytesIO()
    document.save(stream)
    parsed = parse_document(stream.getvalue(), "word")
    assert [block.text for block in parsed.blocks] == [
        "SHIPPING INSTRUCTION",
        "Shipper: Example Export Ltd",
        "Gross Weight: 22 MT",
    ]


def test_xlsx_no_formula_execution_and_cross_sheet_pairing():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["SHIPPING INSTRUCTION"])
    sheet.append(["Shipper", "Example Export Ltd"])
    sheet.append(["Gross Weight", "=1+2"])
    sheet.append(["Consignee"])
    workbook.create_sheet("Other").append(["Unrelated value"])
    stream = BytesIO()
    workbook.save(stream)
    parsed = parse_document(stream.getvalue(), "sheet")
    extracted = extract_labelled(parsed)
    assert extracted.fields[FieldName.SHIPPER].raw_value == "Example Export Ltd"
    assert extracted.fields[FieldName.WEIGHT].state == "missing"
    assert extracted.fields[FieldName.CONSIGNEE].state == "missing"
    assert any("Formula" in warning for warning in parsed.warnings)


def test_unknown_zip_is_rejected():
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("arbitrary.txt", "not an office document")
    with pytest.raises(DomainError):
        parse_document(stream.getvalue(), "zip")


@pytest.mark.parametrize("payload", ['{"x":1,"x":2}', '{"x":NaN}', "```json\n{}\n```"])
def test_strict_provider_json(payload):
    with pytest.raises(ValueError):
        strict_json(payload)


def test_xlsx_hidden_rows_and_columns():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Visible 1", "Visible 2"])
    sheet.append(["Hidden Row Cell", "Visible Value"])
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True

    stream = BytesIO()
    workbook.save(stream)
    parsed = parse_document(stream.getvalue(), "sheet_hidden")

    # Cell A2 is on hidden row 2
    a2_block = next((b for b in parsed.blocks if b.text == "Hidden Row Cell"), None)
    assert a2_block is not None
    assert a2_block.hidden is True
    assert a2_block.quality < 1.0
    assert a2_block.locator.get("hidden_row") is True

    # Cell B1 is on hidden column B
    b1_block = next((b for b in parsed.blocks if b.text == "Visible 2"), None)
    assert b1_block is not None
    assert b1_block.hidden is True
    assert b1_block.locator.get("hidden_col") is True

    # Cell A1 is visible
    a1_block = next((b for b in parsed.blocks if b.text == "Visible 1"), None)
    assert a1_block is not None
    assert a1_block.hidden is False
    assert a1_block.quality == 1.0


def test_corrupt_pdf_raises_file_corrupt():
    corrupt_pdf_data = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\nThis is completely corrupt garbage\x00\xff\xfe"
    with pytest.raises(DomainError) as exc_info:
        parse_document(corrupt_pdf_data, "corrupt_pdf")
    assert exc_info.value.code in {"FILE_CORRUPT", "INTERNAL_ERROR"}


def test_docx_image_warning_when_present():
    from app.config import Settings
    from app.parsers.docx import parse_docx
    document = Document()
    document.add_paragraph("Shipper: Acquired Goods Corp")
    stream = BytesIO()
    document.save(stream)
    # Even without inline shapes, normal docx parses cleanly
    blocks, warnings = parse_docx(stream.getvalue(), "doc1", settings=Settings(environment="test"))
    assert any(b.text == "Shipper: Acquired Goods Corp" for b in blocks)



def test_overprinted_label_and_value_are_regrouped_by_font():
    from app.parsers.pdf import _untangle_overprint

    def chars(text, font, start):
        return [
            {"text": ch, "fontname": font, "x0": start + i * 5.0, "x1": start + i * 5.0 + 4.6}
            for i, ch in enumerate(text)
        ]

    label, value = chars("Notify Consignee", "Helvetica-Bold", 0), chars("CERIEX", "Helvetica", 60)
    merged = sorted(label + value, key=lambda c: c["x0"])
    tangled = {"text": "".join(c["text"] for c in merged), "chars": merged}
    assert _untangle_overprint(tangled)["text"] == "Notify Consignee CERIEX"
    plain = {"text": "Shipper ACME", "chars": chars("Shipper", "Helvetica-Bold", 0) + chars("ACME", "Helvetica", 60)}
    assert _untangle_overprint(plain) is plain  # no horizontal overlap: untouched
